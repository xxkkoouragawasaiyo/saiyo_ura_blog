# A8改善の作業計画（STEP3・4）をExcelにまとめる。元データ：reports/affiliate/2026-10-12/*.md、data/affiliate/link_ledger.csv
# 実行：python3 scripts/build_a8_plan_xlsx.py（その後、数式の再計算が必要）
import re, csv, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

ROOT = '/home/user/saiyo_ura_blog'
OUT = f'{ROOT}/reports/affiliate/2026-10-12/A8改善_作業計画_2026-10-12.xlsx'
FONT = 'Meiryo UI'

plan = open(f'{ROOT}/reports/affiliate/2026-10-12/step3_placement_plan.md', encoding='utf-8').read()
html_blocks = re.findall(r'```html\n(.*?)```', plan, re.S)
sched_md = open(f'{ROOT}/reports/affiliate/2026-10-12/work_schedule.md', encoding='utf-8').read()
html_1479 = re.findall(r'```html\n(.*?)```', sched_md, re.S)[0]

f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_title = Font(name=FONT, size=14, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color='FFFFFF')
f_note = Font(name=FONT, size=9, color='666666')
f_code = Font(name='Consolas', size=9)
fill_head = PatternFill('solid', fgColor='4F6D5E')
fill_input = PatternFill('solid', fgColor='FFF2CC')
fill_p0 = PatternFill('solid', fgColor='F8CBAD')
fill_p1 = PatternFill('solid', fgColor='FCE4D6')
fill_key = PatternFill('solid', fgColor='E2EFDA')
fill_gray = PatternFill('solid', fgColor='EDEDED')
thin = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='top', wrap_text=True)

wb = Workbook()

def setup(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def title(ws, text, sub=None):
    ws['A1'] = text; ws['A1'].font = f_title
    if sub:
        ws['A2'] = sub; ws['A2'].font = f_note

def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = f_head; cell.fill = fill_head; cell.alignment = center; cell.border = border

def put(ws, row, vals, font=f_base, fill=None, align=wrap):
    for i, v in enumerate(vals, 1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = font; cell.alignment = align; cell.border = border
        if fill: cell.fill = fill

# ---------------- はじめに ----------------
ws = wb.active; ws.title = 'はじめに'
setup(ws, [22, 90])
title(ws, 'A8広告の立て直し：作業計画（STEP3・4 確定版）', '作成 2026-10-12／採用のウラ側（saiyo-ura.com）。元データはリポジトリの reports/affiliate/2026-10-12/ の各ファイル。')
rows = [
    ('このファイルの使い方', '上から順にシートを開く。「作業日程」が毎日のToDo。作業が終わったら「状態」を「済」にする。'),
    ('黄色のセル', '入力する欄（状態・実績・日付・数字）。それ以外のセルは触らない（計算式が入っている）。'),
    ('HTMLのコピー', 'HTML_ で始まるシートは、1行＝1セル。B列の指定範囲をまとめて選んでコピー → WordPressのカスタムHTMLブロックに貼る。セルを1つずつダブルクリックしてコピーしない（Excelが「\"」を付けることがある）。'),
    ('', ''),
    ('シート一覧', ''),
    ('作業日程', '日ごとのToDo（優先度・所要時間・使うシート）。1日の合計時間が上限（平日90分・休日120分）を超えていないかを自動で判定。'),
    ('記事別の配置', 'STEP3：上位10記事の「削除・置く・残す」と、添える一文。1行＝1作業。'),
    ('HTML_2112', '2112（インプレピクセル）を空にするHTML。'),
    ('HTML_2111', '2111（パックA・改）の差し替えHTML。'),
    ('HTML_パックJ', '新規パックJ（JAC＋レバテックダイレクト）。979・1489・1827 用。'),
    ('HTML_パックN', '新規パックN（第二新卒エージェントneo）。1790 用。'),
    ('HTML_テキストリンク', '979 用のテキストリンク（レバテックダイレクト・JAC）。'),
    ('1479に貼る文', '10/13 用。2026年度の日程（申込締切 10/24）の表と本文。'),
    ('案件の条件', '使う案件の報酬・成果地点・対象・書けないこと。'),
    ('ルール_AIリライト', 'STEP4 ルールA：AIへの指示文、貼り替え前後の数の比較表（自動判定つき）、作業メモの検索語。'),
    ('ルール_リンク発行', 'STEP4 ルールB〜D：発行手順、A8のサイト番号の判定（a8matを入れると自動で判定）、公開後の確認、置き方。'),
    ('効果測定', 'STEP4：日別の記録表と、基準・2週間・1か月の判定（自動計算）。'),
    ('リンク台帳', '発行したリンクの一覧。新しく発行したら1行足す。'),
    ('別件', '広告以外で見つかった直すこと（優先度つき）。'),
    ('', ''),
    ('重要な日付', '10/13 1479の日程追記（締切 10/24 に間に合わせる）／10/17 基準日（計測の切り替え）／11/1 2週間の判定／11/21 1か月の判定'),
]
r = 4
for k, v in rows:
    ws.cell(row=r, column=1, value=k).font = f_bold
    c = ws.cell(row=r, column=2, value=v); c.font = f_base; c.alignment = wrap
    if k == '黄色のセル': c.fill = fill_input
    r += 1

# ---------------- 作業日程 ----------------
ws = wb.create_sheet('作業日程')
setup(ws, [11, 5, 8, 8, 58, 18, 10, 30])
title(ws, '作業日程（平日90分・休日120分）', '状態（G列）を入力。1日の合計が上限を超えると右の表の「判定」が「超過」になる。日付は2026年。')
header(ws, 4, ['日付', '曜日', '優先', '分', '作業', '使うシート', '状態', 'メモ'])
D = dt.date
tasks = [
 (D(2026,10,13),'P0',30,'1479 に2026年度の日程を追記（申込締切 10/24 に間に合わせる）。IPAのページで日程を最終確認してから貼る','1479に貼る文'),
 (D(2026,10,13),'P0',20,'1107 /aws-saa-difficulty/ の作業メモを削除し、2111 の呼び出しを3か所→1か所に','ルール_AIリライト'),
 (D(2026,10,13),'P1',25,'9/1〜10/12 の日別レポート（A8〔サイト：採用のウラ側〕・VC・もしも・GSC）をダウンロードして共有 → 基準を計算','効果測定'),
 (D(2026,10,13),'',15,'予備',''),
 (D(2026,10,14),'P1',15,'ブログパーツ「パックJ」「パックN」を新規作成。保存後のIDを控えて共有','HTML_パックJ／HTML_パックN'),
 (D(2026,10,14),'P1',30,'979 富士通の配置（パックJ＋テキストリンク2か所。冒頭・途中の2111を削除）','記事別の配置'),
 (D(2026,10,14),'P1',20,'1489 STの配置（パックJ。①の2111と誘導文を削除）','記事別の配置'),
 (D(2026,10,14),'P1',20,'1827 PMの配置（パックJ。冒頭の2111を削除。学習系の一文）','記事別の配置'),
 (D(2026,10,15),'P1',20,'664 応用情報の配置（まとめの2111の前に一文。学習系の一文×2）','記事別の配置'),
 (D(2026,10,15),'P1',20,'1462 DBの配置（冒頭・最後の2111を削除。一文。学習系の一文）','記事別の配置'),
 (D(2026,10,15),'P1',20,'1449 NWの配置（2111を移動。最後の2111を削除）','記事別の配置'),
 (D(2026,10,15),'P1',20,'1143 AWSの配置（2111を移動。最後の2111を削除。学習系の一文）','記事別の配置'),
 (D(2026,10,15),'',10,'予備',''),
 (D(2026,10,16),'P1',20,'1809 PD試験の配置（冒頭の2111を削除。一文。学習系の一文）','記事別の配置'),
 (D(2026,10,16),'P1',15,'1479 SAの配置（最後の2111を削除。①の前に一文）','記事別の配置'),
 (D(2026,10,16),'P1',30,'1790 ITパスの配置（2111×3と誘導文×2を削除、パックN、学習系の一文×2）','記事別の配置'),
 (D(2026,10,16),'P1',25,'10記事をプレビューで見直す（表の崩れ・一文と表の順番）','記事別の配置'),
 (D(2026,10,17),'P1',10,'【基準日】2112を空にする → 2111を差し替える。作業した時刻を記録','HTML_2112／HTML_2111'),
 (D(2026,10,17),'P1',5,'キャッシュを消す（SWELL・ConoHa WING）',''),
 (D(2026,10,17),'P1',20,'公開ページのソースで 0.gif・gifbanner の件数と位置を確認（979・1462・1790）。ボタンは押さない','ルール_リンク発行'),
 (D(2026,10,17),'P1',60,'上位10記事以外で2111を2か所以上呼んでいる記事を探し、1か所に減らす（終わらない分は翌日）','ルール_AIリライト'),
 (D(2026,10,17),'P1',25,'効果測定の記録を始める。リンク台帳の「作業日から」を日付に直す','効果測定／リンク台帳'),
 (D(2026,10,18),'P2',15,'835・892 の TechGO（別サイト001）を、発行済みのサイト002のコードに差し替え','リンク台帳'),
 (D(2026,10,18),'P2',30,'Geekly をサイト002で発行し直し、81・814 を差し替え','ルール_リンク発行'),
 (D(2026,10,18),'P2',15,'787 の UZUZ をサイト002で発行し直して差し替え','ルール_リンク発行'),
 (D(2026,10,18),'P2',10,'81 の検索広告のURL（Google広告のパラメータ付き）を通常のリンクに直す','別件'),
 (D(2026,10,18),'P2',10,'リンク台帳を更新','リンク台帳'),
 (D(2026,10,18),'P2',40,'学習系の案件を検索（A8・もしも・VC：情報処理／基本情報／応用情報／ITパスポート／資格 講座／通信講座／Udemy）→ 一覧を共有',''),
 (D(2026,10,19),'P2',45,'学習系の案件の候補を決めて提携を申請',''),
 (D(2026,10,19),'P2',45,'パックC・Eの「書類選考なし」「書類選考を挟まない」を外す（文言だけ。ピクセル用パーツは触らない）','案件の条件'),
 (D(2026,10,20),'P2',60,'892の「断られた」の見出しなどを確認 → JACの禁止事項に当たるかVCに問い合わせ','別件'),
 (D(2026,10,20),'P2',30,'もしもの type のボタン（2915・2949）に表示計測用画像があるか確認','別件'),
 (D(2026,10,21),'P3',90,'1449 NWの情報を更新（「9月から始めた場合」、シラバス公開済み、経過措置の言い切り）','別件'),
 (D(2026,10,22),'P3',90,'1462 DBの情報を更新（時点、シラバスの公開状況、後期試験の日程）','別件'),
 (D(2026,10,23),'',90,'予備。審査が通った学習系の案件があれば、リンクを発行して一文の後に入れる',''),
 (D(2026,10,24),'P3',120,'1827 PMの更新（後期だけ・申込 1/27〜2/16 の日程、合格率の計算、時点、シラバス）','別件'),
 (D(2026,10,25),'P1',90,'締切後の更新：1489・1449・1479 の「申込は10/24まで」→「申込は終了、試験は11/23まで」＋新試験への案内',''),
 (D(2026,10,25),'P1',30,'効果測定の週次の記録','効果測定'),
 (D(2026,10,26),'P3',30,'1790 の分野の対応表を直す','別件'),
 (D(2026,10,26),'P3',60,'画像の alt（1462・1449・1479）','別件'),
 (D(2026,10,27),'P3',60,'画像の alt（1143・1827・1790）','別件'),
 (D(2026,10,27),'',30,'予備',''),
 (D(2026,10,28),'',90,'予備・学習系の案件の差し込み',''),
 (D(2026,11,1),'P1',30,'【2週間の判定】10/18〜10/31 の日別レポートを記録・共有','効果測定'),
 (D(2026,11,1),'P1',90,'判定の結果で手直し（①が80未満ならソースの点検）','効果測定'),
 (D(2026,11,8),'P1',15,'効果測定の週次の記録','効果測定'),
 (D(2026,11,15),'P1',15,'効果測定の週次の記録','効果測定'),
 (D(2026,11,21),'P1',60,'【1か月の判定】10/18〜11/16 の日別レポートを記録・共有 → 次の手を決める','効果測定'),
 (D(2026,11,22),'P2',90,'パックC（2538/2539）・E（2660/2662）の計測用画像を表の中へ移す（2111と同じ直し方）',''),
]
r = 5
first_row = r
for d, p, m, t, sh in tasks:
    put(ws, r, [d, None, p, m, t, sh, '未着手', None])
    ws.cell(row=r, column=1).number_format = 'm/d'
    ws.cell(row=r, column=2, value=f'=MID("日月火水木金土",WEEKDAY(A{r}),1)').alignment = center
    ws.cell(row=r, column=3).alignment = center
    ws.cell(row=r, column=4).alignment = center
    for col in (7, 8):
        ws.cell(row=r, column=col).fill = fill_input
    if p == 'P0':
        ws.cell(row=r, column=3).fill = fill_p0
    elif p == 'P1':
        ws.cell(row=r, column=3).fill = fill_p1
    r += 1
last_row = r - 1
dv = DataValidation(type='list', formula1='"未着手,作業中,済,延期"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f'G{first_row}:G{last_row}')
ws.conditional_formatting.add(f'A{first_row}:H{last_row}', FormulaRule(formula=[f'$G{first_row}="済"'], fill=fill_gray, font=Font(name=FONT, color='808080')))
ws.freeze_panes = 'A5'

# daily capacity table
ws['J4'] = '日付'; ws['K4'] = '合計（分）'; ws['L4'] = '上限（分）'; ws['M4'] = '判定'
for c in ('J4', 'K4', 'L4', 'M4'):
    ws[c].font = f_head; ws[c].fill = fill_head; ws[c].alignment = center; ws[c].border = border
days = sorted({t[0] for t in tasks})
rr = 5
for d in days:
    ws.cell(row=rr, column=10, value=d).number_format = 'm/d'
    ws.cell(row=rr, column=11, value=f'=SUMIF($A${first_row}:$A${last_row},J{rr},$D${first_row}:$D${last_row})')
    ws.cell(row=rr, column=12, value=f'=IF(WEEKDAY(J{rr},2)>=6,120,90)')
    ws.cell(row=rr, column=13, value=f'=IF(K{rr}>L{rr},"超過","OK")')
    for col in range(10, 14):
        ws.cell(row=rr, column=col).font = f_base; ws.cell(row=rr, column=col).border = border; ws.cell(row=rr, column=col).alignment = center
    rr += 1
ws.cell(row=rr, column=10, value='※ 上限は 平日90分・休日120分（運営者）。祝日は平日扱い。').font = f_note
for col, w in zip('JKLM', [9, 10, 10, 8]):
    ws.column_dimensions[col].width = w

# ---------------- 記事別の配置 ----------------
ws = wb.create_sheet('記事別の配置')
setup(ws, [8, 9, 8, 46, 22, 14, 60, 8])
title(ws, 'STEP3：記事別の配置（確定 2026-10-12）', '区分：削除＝今あるものを消す／置く＝新しく置く／残す＝位置はそのまま。「一文」は表やリンクの直前に入れる段落（体験談は含まない）。')
header(ws, 4, ['記事', '月間クリック', '区分', '位置', '案件', '形式', '添える一文', '済'])
P = [
 ('979', 680, '削除', '導入の吹き出し「IT転職で失敗したくない方へ…」と[2111]①', '', '', ''),
 ('979', 680, '置く', 'H3「【重要】富士通は中途入社のほうが年収が高い」の採用のウラ側メモの直後', 'JAC＋レバテックダイレクト', 'パックJ', 'LevelやSPIRITの確認を、自分から切り出しにくいと感じる方は、エージェントに確認を任せる方法があります。管理・専門職の転職に強いJACと、エンジニアの方なら企業から直接スカウトが届くレバテックダイレクトを比べてみてください。'),
 ('979', 680, '置く', 'H2「課長・部長は何レベル？」最後の採用のウラ側メモの直後（今の[2111]②の位置。表は消す）', 'JAC', 'テキストリンク', 'ポジションの等級をエージェント経由で確認したい方は、【JACリクルートメント】に相談できます。年収600万円以上の方が中心のサービスです。'),
 ('979', 680, '置く', 'H2「まとめ」の吹き出し「IT転職を成功させるためには…」を置き換え（[2111]③は消す）', 'JAC・レバテックダイレクト', 'テキストリンク2行', '富士通への応募を考えている方は、まずLevelとSPIRITの確認から。オファー面談の前に、外の評価も知っておくと判断しやすくなります。＋箇条書き「管理・専門職、年収600万円以上の方 → 【JAC】」「エンジニア・クリエイターの実務経験がある方 → 【レバテックダイレクト】」'),
 ('979', 680, '削除', '[2112]', '', '', ''),
 ('664', 290, '置く', 'H3「『応用情報がなくなるから、今のうちに取るべき』は本当か」の番号付きリストの直後', '学習系（案件決定後）', '一文のみ', '2026年度の試験は、範囲も出題形式も変わりません。今の過去問と参考書で、そのまま準備できます。'),
 ('664', 290, '置く', 'H3「最後の2回はCBT方式…」の「これまでの過去問や参考書が、そのまま使えます。」の直後', '学習系（案件決定後）', 'テキストリンク', '（案件が決まったらリンクだけ）'),
 ('664', 290, '残す', 'H2「まとめ」の後の[2111]（直前に一文を足す）', 'レバテックダイレクト＋TechGo', 'パックA（改）', '応用情報の合格と今の実務経験が、ほかの会社でどう評価されるかを知りたいエンジニアの方は、企業から届くスカウトで確かめられます。'),
 ('664', 290, '削除', '[2112]', '', '', ''),
 ('1462', 180, '削除', '囲み「この記事で分かること」の後の[2111]①', '', '', ''),
 ('1462', 180, '残す', 'H3「新試験…は転職で武器になるか？」のメモの直後の[2111]②（直前に一文）', 'レバテックダイレクト＋TechGo', 'パックA（改）', 'DBの設計やチューニングの経験が、今の職場の外でどう評価されるかは、企業から届くスカウトで確かめられます。'),
 ('1462', 180, '置く', 'H3「2026年度の受験スケジュール」の「今から勉強を始めても十分間に合います。」の直後', '学習系（案件決定後）', '一文のみ', '2026年度のデータベーススペシャリスト試験は、2027年2月頃の実施予定です。'),
 ('1462', 180, '削除', '記事の最後の[2111]③・[2112]', '', '', ''),
 ('1489', 170, '削除', 'H3「新試験『マネジメント領域』の合格は…」の後の誘導文「STの勉強で身につけたことや…」と[2111]①', '', '', ''),
 ('1489', 170, '置く', 'H3「人事から見た答え…」の採用のウラ側メモの後（今の[2111]②の位置。誘導文を差し替え）', 'JAC＋レバテックダイレクト', 'パックJ', 'STの出題範囲と重なる仕事（IT企画・ITコンサル・上流工程）の求人を見てみると、合格の使いどころが具体的になります。管理・専門職に強いJACと、エンジニアの実務経験がある方向けのレバテックダイレクトを比べてみてください。'),
 ('1489', 170, '削除', 'よくある質問の後の[2112]', '', '', ''),
 ('1449', 100, '削除', 'H3「新試験…は転職で武器になるか？」の吹き出しの後の[2111]①', '', '', ''),
 ('1449', 100, '置く', 'H3「ただし『資格名だけ』で評価が上がる時代は終わりつつある」の吹き出しの直後', 'レバテックダイレクト＋TechGo', 'パックA（改）', 'ネットワークの設計・構築の実務経験がある方は、その経験にどんな企業から声がかかるかを、スカウトで確かめられます。'),
 ('1449', 100, '削除', '記事の最後の[2111]②・[2112]', '', '', ''),
 ('1143', 90, '削除', 'H3「『オンプレ経験＋SAP』…」のdodaの段落の後の[2111]①', '', '', ''),
 ('1143', 90, '置く', 'H3「SAPが開くキャリアパス…」の採用のウラ側メモの直後', 'レバテックダイレクト＋TechGo', 'パックA（改）', 'AWSの設計・構築の経験がある方は、クラウドの経験にどんな企業から声がかかるかを、スカウトで確かめられます。'),
 ('1143', 90, '置く', 'H3「おすすめの教材と学習ステップ」のステップ2（Udemyなどの問題集）の段落の後', '学習系（案件決定後）', '一文のみ', '問題集で演習するときは、解説を読んで公式ドキュメントで確かめる使い方が効きます。'),
 ('1143', 90, '削除', '記事の最後の[2111]②・[2112]', '', '', ''),
 ('1809', 80, '削除', '冒頭の※注記の後の[2111]①', '', '', ''),
 ('1809', 80, '置く', 'H3「シラバス案が出た今、できること」の箇条書きの直後', '学習系（案件決定後）', '一文のみ', '受ける領域を決めたら、IPAのシラバス案と並行して、3領域に共通の知識（デジタル戦略・AI基礎・セキュリティ）から固めておくと進めやすくなります。'),
 ('1809', 80, '残す', 'H2「まとめ」の後の[2111]②（直前に一文）', 'レバテックダイレクト＋TechGo', 'パックA（改）', 'すでに実務経験があり、受験と並行して転職も考えている方は、今の経験の評価を先に知っておくと、受ける領域も決めやすくなります。'),
 ('1809', 80, '削除', '[2112]', '', '', ''),
 ('1827', 80, '削除', '囲み「この記事で分かること」の後の[2111]①', '', '', ''),
 ('1827', 80, '置く', 'H2「転職・キャリアで評価されるのか」のメモの後（今の[2111]②の位置）', 'JAC＋レバテックダイレクト', 'パックJ', 'PMとしての経験を活かせる求人は、管理・専門職に強いエージェントに相談すると探しやすくなります。エンジニアとしての実務経験がある方は、企業から直接スカウトを受け取る方法もあります。'),
 ('1827', 80, '削除', '同じ位置の[2112]', '', '', ''),
 ('1827', 80, '置く', 'H2「受験を決めるためのチェックポイント」の「受けると決めたら、次にやること」の直後', '学習系（案件決定後）', '一文のみ', '2026年度のプロジェクトマネージャ試験は、後期（2027年2〜3月）だけの実施です。論文の対策は早めに始めてください。'),
 ('1479', 70, '残す', 'H3「『アーキテクト』を明示的に置く企業が増えている」の最後の[2111]①（直前に一文）', 'レバテックダイレクト＋TechGo', 'パックA（改）', '要件定義や基本設計の経験がある方は、その経験にどんな企業から声がかかるかを、スカウトで確かめられます。'),
 ('1479', 70, '削除', '記事の最後の[2111]②・[2112]', '', '', ''),
 ('1790', 60, '削除', '冒頭の[2111]①', '', '', ''),
 ('1790', 60, '削除', 'H3「迷ったら『受験の目的』で決めましょう」の段落「自分だけで判断が難しいときは、IT業界に詳しい転職エージェントに…」と[2111]②', '', '', ''),
 ('1790', 60, '削除', 'まとめの段落「『自分の場合はどう動くのがいいだろう？』と迷ったら…」と[2111]③・[2112]', '', '', ''),
 ('1790', 60, '置く', 'H3「科目A免除制度は継続される」の箇条書きの直後', '学習系（案件決定後）', '一文のみ', '科目Aの免除を使うには、対象講座の受講と修了試験の合格が必要です。'),
 ('1790', 60, '置く', 'H3「これからの基本情報対策で押さえたい勉強法」の表の直後', '学習系（案件決定後）', '一文のみ', '教材は、2023年以降の出題に対応した最新版を選んでください。'),
 ('1790', 60, '置く', 'H3「なかなか受からないときの考え方」の箇条書きの直後（内部リンク207の前）', '第二新卒エージェントneo', 'パックN', '経歴に自信がない20代の方は、書類づくりから一緒に進めてくれるエージェントに相談する方法もあります。新卒での就職を希望する学生の方は対象外です。'),
]
r = 5; prs = r
for a in P:
    put(ws, r, [a[0], a[1], a[2], a[3], a[4], a[5], a[6], None])
    ws.cell(row=r, column=8).fill = fill_input
    ws.cell(row=r, column=3).alignment = center
    if a[2] == '削除':
        ws.cell(row=r, column=3).font = Font(name=FONT, size=10, color='C00000', bold=True)
    elif a[2] == '置く':
        ws.cell(row=r, column=3).font = Font(name=FONT, size=10, color='1F6B3A', bold=True)
    r += 1
dv2 = DataValidation(type='list', formula1='"済"', allow_blank=True)
ws.add_data_validation(dv2); dv2.add(f'H{prs}:H{r-1}')
ws.conditional_formatting.add(f'A{prs}:H{r-1}', FormulaRule(formula=[f'$H{prs}="済"'], fill=fill_gray))
ws.freeze_panes = 'A5'
r += 1
ws.cell(row=r, column=1, value='表（比較表パック）の数').font = f_bold; r += 1
header(ws, r, ['記事', '変更前', '変更後', '中身', '', '', '', '']); r += 1
for a in [('979',3,1,'パックJ＋テキストリンク2'),('664',1,1,'パックA＋学習系2'),('1462',3,1,'パックA＋学習系1'),('1489',2,1,'パックJ'),('1449',2,1,'パックA'),('1143',2,1,'パックA＋学習系1'),('1809',2,1,'パックA＋学習系1'),('1827',2,1,'パックJ＋学習系1'),('1479',2,1,'パックA'),('1790',3,1,'パックN＋学習系2')]:
    put(ws, r, list(a) + ['', '', '', '']); r += 1
put(ws, r, ['合計', f'=SUM(B{r-10}:B{r-1})', f'=SUM(C{r-10}:C{r-1})', '', '', '', '', ''], font=f_bold)

# ---------------- HTML sheets ----------------
def html_sheet(name, heading, note, code):
    w = wb.create_sheet(name)
    setup(w, [6, 160])
    title(w, heading, note)
    lines = code.rstrip('\n').split('\n')
    start = 5
    w['A4'] = '行'; w['B4'] = f'コード（B{start}〜B{start+len(lines)-1} を選んでコピー）'
    for c in ('A4', 'B4'):
        w[c].font = f_head; w[c].fill = fill_head
    for i, line in enumerate(lines):
        w.cell(row=start+i, column=1, value=i+1).font = f_note
        c = w.cell(row=start+i, column=2)
        c.value = line
        c.data_type = 's'
        c.font = f_code
    w.freeze_panes = 'A5'
    return w

html_sheet('HTML_2112', '2112（パックA_インプレピクセル）を空にする', '10/17（基準日）に、ブログパーツ2112の「︙」→「コードエディター」の中身を全部これに置き換える。', html_blocks[0])
html_sheet('HTML_2111', '2111（パックA_比較表）の差し替え', '10/17（基準日）に、ブログパーツ2111のコードエディターの中身を全部これに置き換える。レバテックダイレクト・TechGo（A8・サイト002）の計測用画像はボタンの直後。', html_blocks[1])
html_sheet('HTML_パックJ', 'パックJ（JAC＋レバテックダイレクト）新規', '10/14 にブログパーツを新規作成して貼る。タイトル例「パックJ_比較表（JAC＋レバテックダイレクト）」。保存後のIDを控える。使う記事：979・1489・1827。', html_blocks[2])
html_sheet('HTML_パックN', 'パックN（第二新卒エージェントneo）新規', '10/14 にブログパーツを新規作成して貼る。タイトル例「パックN_20代（neo）」。保存後のIDを控える。使う記事：1790。', html_blocks[3])

w = wb.create_sheet('HTML_テキストリンク')
setup(w, [24, 150])
title(w, '979 用のテキストリンク', '段落ブロックに貼るとWordPressが画像の行を消すことがあるので、カスタムHTMLブロックで入れる。')
items = [
 ('レバテックダイレクト', html_blocks[4].strip()),
 ('JACリクルートメント', html_blocks[5].strip()),
 ('', ''),
 ('979 課長・部長の章の後', '<p>ポジションの等級をエージェント経由で確認したい方は、' + html_blocks[5].strip() + 'に相談できます。年収600万円以上の方が中心のサービスです。</p>'),
 ('979 まとめ（段落）', '<p>富士通への応募を考えている方は、まずLevelとSPIRITの確認から。オファー面談の前に、外の評価も知っておくと判断しやすくなります。</p>'),
 ('979 まとめ（箇条書き）', '<ul>'),
 ('', '<li>管理・専門職、年収600万円以上の方 → ' + html_blocks[5].strip() + '</li>'),
 ('', '<li>エンジニア・クリエイターの実務経験がある方 → ' + html_blocks[4].strip() + '</li>'),
 ('', '</ul>'),
]
w['A4'] = '用途'; w['B4'] = 'コード（B列のセルを選んでコピー）'
for c in ('A4', 'B4'):
    w[c].font = f_head; w[c].fill = fill_head
for i, (k, v) in enumerate(items):
    w.cell(row=5+i, column=1, value=k).font = f_bold
    c = w.cell(row=5+i, column=2); c.value = v; c.data_type = 's'; c.font = f_code
w.cell(row=5+len(items)+1, column=1, value='まとめの箇条書きは B10〜B13 をまとめて選んでコピー。').font = f_note

html_sheet('1479に貼る文', '1479 に貼る文（10/13）', 'H3「2026年度の受験スケジュール」の、「2026年度は前期試験として2026年11月頃に…」から「…IPA公式ページで必ず最新情報を確認しましょう。」までを置き換える。貼る前にIPAのページで日程を確認する。', html_1479)
w = wb['1479に貼る文']
n0 = 5 + len(html_1479.rstrip('\n').split('\n')) + 1
w.cell(row=n0, column=2, value='ほかに直す3か所（文字だけ書き換える）').font = f_bold
fix = [
 '確定していることの3つ目：「SAは旧春期試験のため、前期試験として2026年11月頃にCBT方式で実施予定」→「SAは旧春期試験のため、2026年度は前期試験としてCBT方式で実施（申込は10月24日まで、科目Aは10月17日から）」',
 'まとめの2つ目：「（前期試験：2026年11月頃・CBT方式）」→「（前期試験：申込は2026年10月24日まで・CBT方式）」',
 '最後の注記：「※この記事の情報は2026年5月時点のものです。」→「※この記事の情報は2026年5月時点のものです（試験日程は、2026年9月17日更新のIPAの案内によります）。」',
 '最後の段落「同じ期に申し込めるのは1区分だけ」は、664の本文（申込は各期1回だけ）による。IPAのページで確認してから貼る。',
]
for i, t in enumerate(fix):
    c = w.cell(row=n0+1+i, column=2, value=t); c.font = f_base; c.alignment = wrap

# ---------------- 案件の条件 ----------------
ws = wb.create_sheet('案件の条件')
setup(ws, [22, 7, 16, 26, 34, 40, 26])
title(ws, '使う案件の条件（2026-10-12 運営者がASPの画面から共有）', '詳しくは program_conditions.md。表やリンクの文言は、この条件とPR文に根拠があることだけを書く。')
header(ws, 4, ['案件', 'ASP', '報酬', '成果地点', '対象', '書けないこと・注意', '使う場所'])
progs = [
 ('レバテックダイレクト', 'A8', '10,000円', '登録＋1日以内の本人確認（最も軽い）', 'エンジニア・クリエイターの実務経験者／50歳未満／国内', '「書類選考なし」の言い切り（根拠は「スカウトの約93%が面談・面接確約」2023年7〜9月）。クリックと登録が同じ端末・ブラウザでないと成果にならない', '2111・パックJ・979テキストリンク'),
 ('TechGo（テックゴー）', 'A8', '18,000円', '30日以内の無料面談実施', '実務2年以上のITエンジニア（20代後半〜30代の支援が中心）／学生は否認', '求人数など根拠のない数字。1Day選考会は事前面談と書類選考の通過が条件。もしも（14,000円）から切り替え', '2111'),
 ('JACリクルートメント', 'VC', '13,384円', '会員登録＋有効判定（軽い）', '年収600万円以上・20代後半〜50代前半／職種は問わない', '年収は「600万〜1500万」「800万〜1500万」のみ。「◯◯に断られた」「登録拒否」NG（サイト内から削除の対象）。取引企業ロゴ・メールのキャプチャーNG', 'パックJ・979テキストリンク'),
 ('第二新卒エージェントneo', 'A8', '15,000円', '30日以内のキャリアカウンセリング', '19〜28歳（18歳以下・29歳以上は否認）／新卒希望の学生・留学生は否認', 'PR文の数字（94.7%・★4.4・113万円UP）は時点がないので使わない。「大手で断られた」の切り口はJACの禁止事項と同じ種類なので使わない', 'パックN（1790）'),
 ('Geekly', 'A8', '来社13,888円／オンライン6,499円', '60日以内の面談完了（重い）', 'IT・WEB・ゲームの経験者／20〜50代／首都圏・関西で働きたい人', '年齢指定・未経験訴求・入稿素材以外の素材・時点のない数字NG', '今回は使わない（81・814はサイト002で発行し直す）'),
 ('レバテックキャリア', 'VC', '11,000円', '登録＋14日以内に本人連絡・エンジニア経験の確認', 'エンジニア経験者／関東・関西・福岡／50代以上は否認', 'レバテックダイレクトと同じ会社。先に登録したほうしか成果にならない → 同じ表・同じ記事に並べない。「No.1」「20万人」は2021年の数字', '892など本人を扱う記事だけ'),
]
r = 5
for p in progs:
    put(ws, r, list(p)); r += 1
ws.freeze_panes = 'A5'

# ---------------- ルール_AIリライト ----------------
ws = wb.create_sheet('ルール_AIリライト')
setup(ws, [12, 12, 9, 9, 9, 9, 9, 9, 9, 9, 9, 9, 9, 9, 12])
title(ws, 'ルールA：AIでリライトするときのチェック', 'A-1 指示文をAIに渡す → A-2 元の本文を保存 → A-3 前後の数を比べる（下の表に入力すると自動判定）→ A-4 作業メモを検索。')
ws['A4'] = 'A-1 AIに渡す指示文（A5〜A9をコピーして、AIへの依頼文に必ず入れる）'; ws['A4'].font = f_bold
prompt = ['【広告と計測に関する決まり（必ず守ること）】',
 '- [blog_parts id="…"] のショートコードと、カスタムHTMLブロック（<!-- wp:html -->）の中身には触れない。移動・削除・追加・書き換えをしない。',
 '- 広告のリンク（px.a8.net／valuecommerce.com／moshimo.com）と、1×1の画像（0.gif／gifbanner／impression）を消さない・分けない・まとめない。',
 '- 作業の説明・確認事項・申し送りは、本文（ブロックの中・HTMLコメントを含む）に書かない。説明は返信の文章で書く。',
 '- 新しい広告の置き場所を提案する場合は、本文に入れず、「どの見出しの後に・何を」を返信の文章で書く。']
for i, t in enumerate(prompt):
    c = ws.cell(row=5+i, column=1, value=t); c.font = f_code
ws['A11'] = 'A-2 貼り替え前に：コードエディターの中身を全部保存（ファイル名：投稿ID_日付_before.txt）。リビジョンの件数を控える。'; ws['A11'].font = f_bold
ws['A13'] = 'A-3 前後の数の比較（コードエディターの中身で Ctrl+F の件数を入力。黄色が入力欄。判定は自動）'; ws['A13'].font = f_bold
cols = ['投稿ID', '日付', '[blog_parts id= 前', '後', 'px.a8.net 前', '後', '0.gif 前', '後', 'valuecommerce 前', '後', 'i.moshimo 前', '後', 'post-link 前', '後', '判定']
header(ws, 14, cols)
ex = ['（例）1107', dt.date(2026,10,13), 3, 1, 2, 2, 2, 2, 0, 0, 0, 0, 3, 3]
for i, v in enumerate(ex, 1):
    c = ws.cell(row=15, column=i, value=v); c.font = f_note; c.border = border
ws.cell(row=15, column=2).number_format = 'm/d'
for rr in range(15, 41):
    if rr > 15:
        for col in range(1, 15):
            c = ws.cell(row=rr, column=col); c.fill = fill_input; c.border = border; c.font = f_base
        ws.cell(row=rr, column=2).number_format = 'm/d'
    ws.cell(row=rr, column=15, value=(
        f'=IF(A{rr}="","",IF(AND(E{rr}=F{rr},G{rr}=H{rr},I{rr}=J{rr},K{rr}=L{rr},N{rr}>=M{rr},F{rr}=H{rr}),'
        f'IF(C{rr}=D{rr},"OK","OK（表の数を変えた：意図どおりか確認）"),"要確認：数が合わない")'))
    ws.cell(row=rr, column=15).font = f_bold; ws.cell(row=rr, column=15).border = border
ws.column_dimensions['O'].width = 34
ws['A42'] = '判定の決まり：A8のリンクと0.gifの数が同じ（F=H）、各リンク・計測用画像の数が前後で同じ、内部リンクカードが減っていない。表の呼び出し（[blog_parts）は、意図して減らしたときだけ違ってよい。15行目は記入例。'
ws['A42'].font = f_note
ws['A44'] = 'A-4 公開前に検索して0件であること（作業メモの残り）'; ws['A44'].font = f_bold
words = ['要確認', '差し替え', '削除済み', '指定どおり', '以下の', 'TODO', '000X', 'AI', 'Claude', 'ChatGPT', '<!-- （<!-- wp: と <!-- /wp: 以外）']
for i, t in enumerate(words):
    ws.cell(row=45+i, column=1, value=t).font = f_base
ws['A57'] = 'A-5 広告の位置が変わった記事は「ルール_リンク発行」のルールD（置き方）を確認。 A-6 IPA記事は本文の時点・日程・シラバスの記載を更新（664・1489・1809に合わせる）。'
ws['A57'].font = f_note

# ---------------- ルール_リンク発行 ----------------
ws = wb.create_sheet('ルール_リンク発行')
setup(ws, [16, 20, 16, 24, 60])
title(ws, 'ルールB〜D：リンクの発行・公開後の確認・置き方', 'B列（黄色）に a8mat= の2つ目のかたまりを入れると、サイト番号を自動で判定する。')
ws['A4'] = 'B-2 A8のサイト番号の判定'; ws['A4'].font = f_bold
header(ws, 5, ['案件・メモ', 'a8matの2つ目', '10進数', 'サイト番号', '判定'])
checks = [('レバテックダイレクト', '786GWQ'), ('neo', '72TKGQ'), ('TechGo', 'F172FU'), ('814 の Geekly（例）', 'EZEQUQ')] + [('', '')] * 6
CH = '"0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"'
for i, (k, v) in enumerate(checks):
    rr = 6 + i
    ws.cell(row=rr, column=1, value=k).font = f_base
    c = ws.cell(row=rr, column=2, value=v if v else None); c.fill = fill_input; c.font = f_base
    ws.cell(row=rr, column=3, value=(
        f'=IF(B{rr}="","",SUMPRODUCT((FIND(MID(UPPER(B{rr}),ROW(INDIRECT("1:"&LEN(B{rr}))),1),{CH})-1)'
        f'*36^(LEN(B{rr})-ROW(INDIRECT("1:"&LEN(B{rr}))))))')).number_format = '0'
    ws.cell(row=rr, column=4, value=f'=IF(B{rr}="","",MID(TEXT(C{rr},"000000000"),4,3))')
    ws.cell(row=rr, column=5, value=f'=IF(B{rr}="","",IF(D{rr}="002","採用のウラ側（OK）","別サイト → 発行し直す"))')
    for col in range(1, 6):
        ws.cell(row=rr, column=col).border = border
        if col > 2: ws.cell(row=rr, column=col).font = f_base
ws['A16'] = '例：a8mat=4B1ILL+72TKGQ+3Y6M+63H8I なら、2つ目は 72TKGQ。Excelでは =MID(TEXT(DECIMAL("72TKGQ",36),"000000000"),4,3) でも同じ結果になる。'
ws['A16'].font = f_note
r = 18
blocks = [
 ('B-1 発行する前', ['案件の掲載条件（成果条件・否認条件・禁止事項・PR文・報酬）を記録する（「案件の条件」シート）。', '置く記事の読者が対象条件に入るか確かめる（年齢・経験・地域・学生かどうか）。', '禁止表現を確かめる（Geekly：年齢指定・未経験訴求NG／JAC：年収は2パターンのみ・「断られた」NG）。']),
 ('B-2 発行するとき', ['A8：画面上部のサイト選択が「採用のウラ側」。上の表でサイト番号が 002。', 'VC：「現在のサイト：採用のウラ側（ID:3765202）」。コードが sid=3765202。', 'もしも：サイト選択が「採用のウラ側」（でんチャレではない）。コードが a_id=5422590。']),
 ('B-3 発行したあと', ['発行したコードはそのまま使う。0.gif・gifbanner・impression の行を消さない・リンクと分けない。', 'カスタムHTMLブロックに貼る（段落ブロックだと画像の行が消えることがある）。', '計測用画像はボタン（リンク）の直後。別パーツや記事の最後に分けない。画面の外（left:-9999px など）に置かない。', '「リンク台帳」に1行足す。']),
 ('B-4', ['同じ会社の案件（レバテックダイレクトとレバテックキャリアなど）は同じ表・同じ記事に並べない。']),
 ('ルールC 公開後', ['公開ページのソース（Ctrl+U）で 0.gif・gifbanner・i.moshimo.com を検索。件数が「表＋テキストリンクの数」と合うこと、ボタンの直後にあること（data-src でもよい）。', '表のボタンは押さない。リンク先はASPの管理画面の「リンク先を見る」で確かめる。', '確認で自分が表示・クリックした回数を「効果測定」シートに記録する。']),
 ('ルールD 置き方', ['比較表は1記事1か所（長い記事でも2か所まで）。', '冒頭（最初のH2の前）に表を置かない。', '表の直前に誘導の一文。一文は案件の形（スカウト型／エージェント）と合わせる。', '読者と案件の対象条件を照らし合わせる。', '説明文は掲載条件・PR文に根拠があることだけ。数字には時点。「No.1」「書類選考なし」は使わない。', '新しいパックのCSS（.saiyo-cta）は既存と同じにする。']),
]
for h, items_ in blocks:
    ws.cell(row=r, column=1, value=h).font = f_bold; r += 1
    for t in items_:
        c = ws.cell(row=r, column=2, value='・' + t); c.font = f_base
        r += 1
    r += 1

# ---------------- 効果測定 ----------------
ws = wb.create_sheet('効果測定')
setup(ws, [11, 11, 11, 9, 9, 9, 9, 9, 9, 9, 9, 9, 13, 12, 26, 13])
title(ws, '効果測定（STEP4）', '黄色に入力。ASPの数字は「採用のウラ側」で絞り込んだ日別の値。流入はGSCの検索クリック。判定は自動計算。')
ws['A4'] = '期間（黄色は変更可）'; ws['A4'].font = f_bold
periods = [('基準', dt.date(2026,9,1), dt.date(2026,10,13)), ('移行（除外）', dt.date(2026,10,14), dt.date(2026,10,17)), ('2週間', dt.date(2026,10,18), dt.date(2026,10,31)), ('1か月', dt.date(2026,10,18), dt.date(2026,11,16))]
header(ws, 5, ['期間', '開始', '終了'] + [''] * 13)
for i, (k, s_, e_) in enumerate(periods):
    rr = 6 + i
    ws.cell(row=rr, column=1, value=k).font = f_bold
    for col, v in ((2, s_), (3, e_)):
        c = ws.cell(row=rr, column=col, value=v); c.number_format = 'yyyy/m/d'; c.fill = fill_input; c.font = f_base; c.border = border
ws['E6'] = '基準は9/1〜配置変更の前日。7月（自分での表示）と8月（計測の形が途中で変わった）は使わない。基準日は10/17（2112を空にし2111を差し替えた日）。'
ws['E6'].font = f_note

# summary
ws['A11'] = '判定'; ws['A11'].font = f_bold
header(ws, 12, ['指標', '基準', '2週間', '1か月', '2週間の成功', '2週間の失敗', '1か月の成功', '2週間の判定', '1か月の判定', '', '', '', '', '', '', ''])
DR = 31; DE = DR + 120  # daily rows range
def S(col, k):
    s_ = f'$B${6+k}'; e_ = f'$C${6+k}'
    return f'SUMIFS({col}${DR}:{col}${DE},$A${DR}:$A${DE},">="&{s_},$A${DR}:$A${DE},"<="&{e_})'
# helper rows 13..: sums
labels = ['流入（GSCクリック）', 'A8表示', 'A8クリック', '全ASPクリック', '上位10記事の案件の発生']
colmap = {'流入（GSCクリック）': 'B', 'A8表示': 'D', 'A8クリック': 'E'}
rr = 13
sum_rows = {}
for lab in labels:
    ws.cell(row=rr, column=1, value=lab).font = f_base
    for j, k in ((2, 0), (3, 2), (4, 3)):
        if lab == '全ASPクリック':
            f = f'={S("E",k)}+{S("H",k)}+{S("K",k)}'
        elif lab == '上位10記事の案件の発生':
            f = f'={S("N",k)}'
        else:
            f = f'={S(colmap[lab],k)}'
        ws.cell(row=rr, column=j, value=f).number_format = '#,##0'
    sum_rows[lab] = rr
    for col in range(1, 10): ws.cell(row=rr, column=col).border = border
    rr += 1
g = sum_rows['流入（GSCクリック）']; a8i = sum_rows['A8表示']; a8c = sum_rows['A8クリック']; allc = sum_rows['全ASPクリック']; occ = sum_rows['上位10記事の案件の発生']
metrics = [
 ('① A8表示／流入1,000', a8i, '0.0', '>=120', '<80', '>=120'),
 ('② A8クリック／流入1,000', a8c, '0.00', '>=1.2', '<0.5', '>=2'),
 ('③ 全ASPクリック／流入1,000', allc, '0.00', 'base+1', None, 'base+2'),
 ('④ A8クリック率', None, '0.00%', None, None, '>=0.008'),
 ('⑤ 上位10記事の案件の発生（件）', occ, '0', None, None, '>=1'),
]
for lab, src, fmt, s2, f2, s1 in metrics:
    ws.cell(row=rr, column=1, value=lab).font = f_bold
    for j, colL in ((2, 'B'), (3, 'C'), (4, 'D')):
        if lab.startswith('④'):
            f = f'=IF({colL}{a8i}=0,"",{colL}{a8c}/{colL}{a8i})'
        elif lab.startswith('⑤'):
            f = f'={colL}{src}'
        else:
            f = f'=IF({colL}{g}=0,"",{colL}{src}/{colL}{g}*1000)'
        c = ws.cell(row=rr, column=j, value=f); c.number_format = fmt; c.fill = fill_key
    # thresholds text & judgments
    if lab.startswith('①'):
        ws.cell(row=rr, column=5, value=120); ws.cell(row=rr, column=6, value=80); ws.cell(row=rr, column=7, value='120以上を保つ')
        ws.cell(row=rr, column=8, value=f'=IF(C{rr}="","未入力",IF(C{rr}>=E{rr},"成功（計測は直った）",IF(C{rr}<F{rr},"失敗：ソースを点検","途中：記事ごとに確認")))')
        ws.cell(row=rr, column=9, value=f'=IF(D{rr}="","未入力",IF(D{rr}>=E{rr},"OK","要確認"))')
    elif lab.startswith('②'):
        ws.cell(row=rr, column=5, value=1.2); ws.cell(row=rr, column=6, value=0.5); ws.cell(row=rr, column=7, value=2.0)
        ws.cell(row=rr, column=8, value=f'=IF(C{rr}="","未入力",IF(C{rr}>=E{rr},"順調（参考）",IF(C{rr}<F{rr},"要注意（参考）","途中（参考）")))')
        ws.cell(row=rr, column=9, value=f'=IF(D{rr}="","未入力",IF(D{rr}>=G{rr},"成功","未達：訴求・位置を見直す"))')
    elif lab.startswith('③'):
        ws.cell(row=rr, column=5, value=f'=IF(B{rr}="","",B{rr}+1)'); ws.cell(row=rr, column=6, value=f'=IF(B{rr}="","",B{rr})'); ws.cell(row=rr, column=7, value=f'=IF(B{rr}="","",B{rr}+2)')
        for col in (5, 6, 7): ws.cell(row=rr, column=col).number_format = '0.00'
        ws.cell(row=rr, column=8, value=f'=IF(OR(B{rr}="",C{rr}=""),"未入力",IF(C{rr}>=E{rr},"順調（参考）",IF(C{rr}<F{rr},"基準より下：D・Eの記事を確認","途中（参考）")))')
        ws.cell(row=rr, column=9, value=f'=IF(OR(B{rr}="",D{rr}=""),"未入力",IF(D{rr}>=G{rr},"成功","未達"))')
    elif lab.startswith('④'):
        ws.cell(row=rr, column=7, value=0.008).number_format = '0.0%'
        ws.cell(row=rr, column=9, value=f'=IF(D{rr}="","未入力",IF(D{rr}>=G{rr},"OK","表示は増えたが押されていない"))')
    elif lab.startswith('⑤'):
        ws.cell(row=rr, column=7, value=1)
        ws.cell(row=rr, column=9, value=f'=IF(D{g}=0,"未入力",IF(D{rr}>=G{rr},"成功","0件：クリックの中身を確認"))')
    for col in range(1, 10):
        ws.cell(row=rr, column=col).border = border
        if col >= 5: ws.cell(row=rr, column=col).font = f_base
    rr += 1
ws.cell(row=rr, column=1, value='③は「基準＋1.0（2週間・参考）」「基準＋2.0（1か月）」。①の基準の裏づけ：計測が正しく働いているVCの表は流入1,000あたり約153表示（9/10〜10/10）。④の判定は1か月の値で見る。').font = f_note
ws.cell(row=rr+1, column=1, value='見込みの前提：A8の計測用画像が入った表のある記事に流入の30〜40%、表まで読む人30〜50%、表1つの画像1〜2個、表のクリック率2〜4%（業界の一般的な水準で、当サイトの実測ではない）。①が120以上になったら、その2週間の実測を新しい基準にする。').font = f_note

# reference table
ws.cell(row=rr+3, column=1, value='参考：今の数字（運営者の共有・リポジトリのデータ）').font = f_bold
ref = [('2026-09 A8', '97表示・2クリック', '流入約4,230（141／日）', '①約23・②約0.5'),
       ('2026-10（1〜12日）A8', '64表示・2クリック', '流入約2,400（約200／日）', '①約27・②約0.8'),
       ('9/10〜10/10 VC', '646表示・22クリック・1件（成果報酬330円・案件は要確認）', 'GSC 27日分 4,230', '表示約153・クリック約5.2'),
       ('全ASP（9/10〜10/10）', '約26クリック', '', '③約6.2')]
for i, rw in enumerate(ref):
    for j, v in enumerate(rw, 1):
        c = ws.cell(row=rr+4+i, column=j, value=v); c.font = f_note

# daily table
ws.cell(row=DR-2, column=1, value='日別の記録（黄色に入力。週1回まとめて入力してもよい）').font = f_bold
cols = ['日付', 'GSC（全体）', 'GSC（上位10記事）', 'A8表示', 'A8クリック', 'A8発生', 'VC表示', 'VCクリック', 'VC発生', 'もしも表示', 'もしもクリック', 'もしも発生', '自分で見た・押した', '上位10記事の案件の発生', 'メモ', '期間']
header(ws, DR-1, cols)
d0 = dt.date(2026, 9, 1)
for i in range(DE - DR + 1):
    rr = DR + i
    d = d0 + dt.timedelta(days=i)
    c = ws.cell(row=rr, column=1, value=d); c.number_format = 'm/d(aaa)'; c.font = f_base; c.border = border
    for col in range(2, 16):
        cc = ws.cell(row=rr, column=col); cc.fill = fill_input; cc.border = border; cc.font = f_base
    ws.cell(row=rr, column=16, value=(f'=IF(AND(A{rr}>=$B$6,A{rr}<=$C$6),"基準",IF(AND(A{rr}>=$B$7,A{rr}<=$C$7),"移行（除外）",'
                                       f'IF(AND(A{rr}>=$B$8,A{rr}<=$C$8),"2週間・1か月",IF(AND(A{rr}>=$B$9,A{rr}<=$C$9),"1か月",""))))')).font = f_note
    if d == dt.date(2026, 10, 17):
        ws.cell(row=rr, column=15, value='基準日（2112を空にし2111を差し替え）')
# no freeze (table starts low on the sheet)

# ---------------- リンク台帳 ----------------
ws = wb.create_sheet('リンク台帳')
setup(ws, [20, 7, 16, 30, 38, 26, 12, 26, 16, 40])
title(ws, 'リンク台帳', '新しく発行したら1行足す（黄色）。別サイト（001）のものは「要差し替え」。')
rows_ = list(csv.reader(open(f'{ROOT}/data/affiliate/link_ledger.csv', encoding='utf-8')))
header(ws, 4, rows_[0])
r = 5
for row in rows_[1:]:
    put(ws, r, row)
    if row[8] == '要差し替え':
        ws.cell(row=r, column=9).fill = fill_p0
    r += 1
for rr in range(r, r + 8):
    for col in range(1, 11):
        c = ws.cell(row=rr, column=col); c.fill = fill_input; c.border = border; c.font = f_base
ws.freeze_panes = 'A5'

# ---------------- 別件 ----------------
ws = wb.create_sheet('別件')
setup(ws, [7, 70, 22, 12])
title(ws, '広告以外で見つかった直すこと', '「作業日程」に入っているものは日付を記載。状態を入力。')
header(ws, 4, ['優先', '内容', '予定', '状態'])
items_ = [
 ('高', '1107 /aws-saa-difficulty/ の本文に残っているAIの作業メモを削除（2111を1か所に）', '10/13'),
 ('高', '1479 に2026年度の日程（申込締切 10/24）を追記', '10/13'),
 ('高', 'ST・NW・SA（1489・1449・1479）の締切後の更新', '10/25'),
 ('中', '別サイト（001）のコードを差し替え：81・814 Geekly、835・892 TechGO、787 UZUZ', '10/18'),
 ('中', '81 の検索広告のURL（Google広告のパラメータ付き）を通常のリンクに直す', '10/18'),
 ('中', 'パックC・Eの「書類選考なし」「書類選考を挟まない」を外す', '10/19'),
 ('中', '892 の「断られた」など、JACの禁止事項に当たるかVCに確認', '10/20'),
 ('中', 'パックC（2538/2539）・E（2660/2662）の計測用画像を表の中へ', '11/22（1か月の判定後）'),
 ('低', 'もしもの type のボタン（2915・2949）に表示計測用画像があるか確認', '10/20'),
 ('低', '1449 の古い記述（「9月から始めた場合」、シラバス、経過措置の言い切り）', '10/21'),
 ('低', '1462 の古い記述（時点、シラバス、後期試験の日程）', '10/22'),
 ('低', '1827：2026年度の日程がない、合格率の計算（1,066÷12,197は約8.7%で13.5%にならない）、経過措置の言い切り、時点', '10/24'),
 ('低', '1790：ITパスポートの分野の対応表が1対1に見える、時点の記載がない', '10/26'),
 ('低', '979：同じ文が2回（「転職時のオファー年収は…」）、「2026年」の「2」が太字の外', '未定'),
 ('低', '664：まとめの太字が二重。1809・1449 の最後に空の段落', '未定'),
 ('低', '画像の alt が空（1462・1449・1143・1827・1479・1790 ほか）', '10/26〜27'),
 ('低', '1143 の「この記事の結論まとめ」が画像1枚だけ（文字がない）', '未定'),
 ('低', 'VCで9/10〜10/10に発生した1件（成果報酬330円）がどの案件か確認', '未定'),
]
r = 5
for p, t, d in items_:
    put(ws, r, [p, t, d, '未着手'])
    ws.cell(row=r, column=4).fill = fill_input
    if p == '高': ws.cell(row=r, column=1).fill = fill_p0
    r += 1
dv3 = DataValidation(type='list', formula1='"未着手,作業中,済,延期"', allow_blank=True)
ws.add_data_validation(dv3); dv3.add(f'D5:D{r-1}')

for w in wb.worksheets:
    w.sheet_view.zoomScale = 100
wb.save(OUT)
print(OUT)
